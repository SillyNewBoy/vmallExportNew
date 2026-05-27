from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"

FONT = "Microsoft YaHei"

# =============================================================================
# 使用前只需要修改这里
# =============================================================================
#
# 1. PRODUCT_NAME：报告品类名，例如 "手机"、"平板"、"耳机"。
# 2. INPUT_XLSX：人工评测结果 Excel 路径。
# 3. OUTPUT_PPTX：PPT 输出路径。
# 4. 如果 PPT 正在被打开，Windows 会禁止覆盖同名文件；请关闭 PPT 或修改 OUTPUT_PPTX。
#
CONFIG = {
    "PRODUCT_NAME": "手机",
    "INPUT_XLSX": REPORT_DIR / "手机全量参数评测.xlsx",
    "OUTPUT_PPTX": ROOT / "小艺手机全量参数评测报告.pptx",
}

COLORS = {
    "red": "C7000B",
    "red_dark": "8F0008",
    "dark": "1F2329",
    "gray": "6B7280",
    "light_gray": "EEF1F5",
    "mid_gray": "D5DAE1",
    "bg": "F7F8FA",
    "white": "FFFFFF",
    "green": "16845B",
    "blue": "2563EB",
    "orange": "D97706",
    "amber": "F59E0B",
    "purple": "7C3AED",
}

DIMENSIONS = ["相关性", "正确性", "信息量", "一致性"]
BASE_COLUMNS = [
    "case_id",
    "四级分类",
    "query所属设备",
    "五级分类",
    "六级分类",
    "query",
    "DisplayText",
    "备注",
    "人工综合评测结果",
    *DIMENSIONS,
]

def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def truncate(text: object, limit: int) -> str:
    value = "" if text is None else str(text)
    return value if len(value) <= limit else value[: limit - 1] + "…"


def fitted_font_size(text: object, base_size: float, *, fit_chars: int, min_size: float = 5.5) -> float:
    value = "" if text is None else str(text)
    if len(value) <= fit_chars:
        return base_size
    shrink = (len(value) - fit_chars + 5) // 6
    return max(min_size, base_size - shrink)


def load_eval_rows(input_path: Path) -> list[dict[str, object]]:
    wb = load_workbook(input_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(value).strip() if value is not None else "" for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    index = {header: idx for idx, header in enumerate(headers)}
    missing = [column for column in BASE_COLUMNS if column not in index]
    if missing:
        raise ValueError(f"评测表缺少必要字段：{', '.join(missing)}")

    rows: list[dict[str, object]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        row = {column: values[index[column]] for column in BASE_COLUMNS}
        row["case_id"] = str(row["case_id"] or "").strip()
        row["四级分类"] = str(row["四级分类"] or "").strip()
        row["query所属设备"] = str(row["query所属设备"] or "未标注设备").strip()
        row["五级分类"] = str(row["五级分类"] or "未标注五级分类").strip()
        row["六级分类"] = str(row["六级分类"] or "未标注六级分类").strip()
        row["query"] = str(row["query"] or "").strip()
        row["DisplayText"] = str(row["DisplayText"] or "").strip()
        row["备注"] = str(row["备注"] or "").strip()
        for column in ["人工综合评测结果", *DIMENSIONS]:
            row[column] = parse_flag(row[column], column)
        rows.append(row)
    if not rows:
        raise ValueError("评测表没有可用数据行")
    return rows


def parse_flag(value: object, column: str) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return 1 if int(value) == 1 else 0
    text = str(value).strip()
    if text in {"1", "是", "有问题", "问题", "true", "TRUE"}:
        return 1
    if text in {"0", "否", "无问题", "通过", "false", "FALSE"}:
        return 0
    raise ValueError(f"{column} 字段只支持 0/1 或等价文本，发现：{value!r}")


def compute_summary(rows: list[dict[str, object]]) -> dict[str, object]:
    total = len(rows)
    issue_count = sum(int(row["人工综合评测结果"]) for row in rows)
    pass_count = total - issue_count
    dim_counts = {dim: sum(int(row[dim]) for row in rows) for dim in DIMENSIONS}

    by_l5: dict[str, dict[str, object]] = {}
    for l5 in sorted({str(row["五级分类"]) for row in rows}):
        subset = [row for row in rows if row["五级分类"] == l5]
        issue = sum(int(row["人工综合评测结果"]) for row in subset)
        by_l5[l5] = {
            "total": len(subset),
            "issue": issue,
            "pass": len(subset) - issue,
            "accuracy": (len(subset) - issue) / len(subset),
            "issue_rate": issue / len(subset),
            "dims": {dim: sum(int(row[dim]) for row in subset) for dim in DIMENSIONS},
        }

    by_l6: dict[tuple[str, str], dict[str, object]] = {}
    groups: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[(str(row["五级分类"]), str(row["六级分类"]))].append(row)
    for key, subset in groups.items():
        issue = sum(int(row["人工综合评测结果"]) for row in subset)
        by_l6[key] = {
            "total": len(subset),
            "issue": issue,
            "pass": len(subset) - issue,
            "accuracy": (len(subset) - issue) / len(subset),
            "issue_rate": issue / len(subset),
            "dims": {dim: sum(int(row[dim]) for row in subset) for dim in DIMENSIONS},
        }

    by_l6_device: dict[tuple[str, str], list[tuple[str, int]]] = {}
    for key, subset in groups.items():
        device_counter: Counter[str] = Counter()
        for row in subset:
            if int(row["人工综合评测结果"]) == 1:
                device_counter[str(row["query所属设备"])] += 1
        by_l6_device[key] = device_counter.most_common(3)

    by_device: dict[str, dict[str, object]] = {}
    for device in sorted({str(row["query所属设备"]) for row in rows}):
        subset = [row for row in rows if row["query所属设备"] == device]
        issue = sum(int(row["人工综合评测结果"]) for row in subset)
        by_device[device] = {
            "total": len(subset),
            "issue": issue,
            "pass": len(subset) - issue,
            "accuracy": (len(subset) - issue) / len(subset),
            "issue_rate": issue / len(subset),
            "dims": {dim: sum(int(row[dim]) for row in subset) for dim in DIMENSIONS},
        }

    examples = []
    seen_l5: set[str] = set()
    issue_rows = [row for row in rows if int(row["人工综合评测结果"]) == 1]
    issue_rows.sort(
        key=lambda row: (
            int(row["正确性"]),
            int(row["信息量"]),
            int(row["相关性"]),
            int(row["一致性"]),
        ),
        reverse=True,
    )
    for row in issue_rows:
        l5 = str(row["五级分类"])
        if l5 in seen_l5 and len(examples) < 4:
            continue
        examples.append(row)
        seen_l5.add(l5)
        if len(examples) == 4:
            break

    top_l6 = sorted(by_l6.items(), key=lambda item: item[1]["issue"], reverse=True)
    device_rank = sorted(by_device.items(), key=lambda item: item[1]["accuracy"])
    top_dim = max(dim_counts.items(), key=lambda item: item[1])[0]
    worst_l5 = min(by_l5.items(), key=lambda item: item[1]["accuracy"])[0]
    worst_device = device_rank[0][0]
    worst_dim_device = find_worst_dim_device(by_device, top_dim)
    device_bands = build_device_bands(device_rank)
    bad_cases = build_bad_cases(rows, device_rank, by_l5)
    low_accuracy_failures = build_low_accuracy_failures(rows, by_device)

    return {
        "total": total,
        "issue_count": issue_count,
        "pass_count": pass_count,
        "accuracy": pass_count / total,
        "issue_rate": issue_count / total,
        "dim_counts": dim_counts,
        "by_l5": by_l5,
        "by_l6": by_l6,
        "by_l6_device": by_l6_device,
        "by_device": by_device,
        "device_rank": device_rank,
        "top_l6": top_l6,
        "examples": examples,
        "top_dim": top_dim,
        "worst_l5": worst_l5,
        "worst_device": worst_device,
        "worst_dim_device": worst_dim_device,
        "device_bands": device_bands,
        "bad_cases": bad_cases,
        "low_accuracy_failures": low_accuracy_failures,
    }


def find_worst_dim_device(by_device: dict[str, dict[str, object]], dim: str) -> dict[str, object]:
    ranked = sorted(
        by_device.items(),
        key=lambda item: (
            item[1]["dims"][dim] / item[1]["total"],
            item[1]["dims"][dim],
            item[1]["issue_rate"],
        ),
        reverse=True,
    )
    device, stat = ranked[0]
    return {
        "device": device,
        "dimension": dim,
        "dimension_issue": stat["dims"][dim],
        "dimension_rate": stat["dims"][dim] / stat["total"],
        "total": stat["total"],
        "accuracy": stat["accuracy"],
    }


def build_device_bands(device_rank: list[tuple[str, dict[str, object]]]) -> dict[str, list[tuple[str, dict[str, object]]]]:
    return {
        "90%-100%": [(device, stat) for device, stat in device_rank if 0.90 < stat["accuracy"] <= 1.00],
        "60%-90%": [(device, stat) for device, stat in device_rank if 0.60 <= stat["accuracy"] <= 0.90],
        "30%-60%": [(device, stat) for device, stat in device_rank if 0.30 < stat["accuracy"] < 0.60],
        "0%-30%": [(device, stat) for device, stat in device_rank if 0 <= stat["accuracy"] <= 0.30],
    }


def build_bad_cases(
    rows: list[dict[str, object]],
    device_rank: list[tuple[str, dict[str, object]]],
    by_l5: dict[str, dict[str, object]],
) -> dict[str, list[dict[str, object]]]:
    issue_rows = [row for row in rows if int(row["人工综合评测结果"]) == 1]
    worst_devices = [device for device, _ in device_rank[:2]]
    worst_l5s = [l5 for l5, _ in sorted(by_l5.items(), key=lambda item: item[1]["accuracy"])[:2]]
    device_cases = pick_issue_rows(issue_rows, "query所属设备", worst_devices, limit=2)
    used_case_ids = {str(row["case_id"]) for row in device_cases}
    return {
        "devices": device_cases,
        "l5": pick_issue_rows(issue_rows, "五级分类", worst_l5s, limit=2, exclude_case_ids=used_case_ids),
    }


def pick_issue_rows(
    issue_rows: list[dict[str, object]],
    field: str,
    values: list[str],
    *,
    limit: int,
    exclude_case_ids: set[str] | None = None,
) -> list[dict[str, object]]:
    picked: list[dict[str, object]] = []
    seen: set[str] = set(exclude_case_ids or set())
    for value in values:
        for row in issue_rows:
            case_id = str(row["case_id"])
            if row[field] == value and case_id not in seen:
                picked.append(row)
                seen.add(case_id)
                break
        if len(picked) >= limit:
            break
    if len(picked) < limit:
        for row in issue_rows:
            case_id = str(row["case_id"])
            if case_id in seen:
                continue
            picked.append(row)
            seen.add(case_id)
            if len(picked) >= limit:
                break
    return picked


def build_low_accuracy_failures(
    rows: list[dict[str, object]],
    by_device: dict[str, dict[str, object]],
) -> list[dict[str, object]]:
    result: list[dict[str, object]] = []
    for device, stat in sorted(by_device.items(), key=lambda item: item[1]["accuracy"]):
        if stat["accuracy"] > 0.30:
            continue
        issue_rows = [
            row
            for row in rows
            if row["query所属设备"] == device and int(row["人工综合评测结果"]) == 1
        ]
        params = Counter(f"{row['五级分类']} / {row['六级分类']}" for row in issue_rows)
        dim_counts = {dim: sum(int(row[dim]) for row in issue_rows) for dim in DIMENSIONS}
        result.append(
            {
                "device": device,
                "total": stat["total"],
                "accuracy": stat["accuracy"],
                "issue": stat["issue"],
                "params": params.most_common(),
                "top_dim": max(dim_counts.items(), key=lambda item: item[1])[0] if issue_rows else "-",
            }
        )
    return result


def write_summary_workbook(summary: dict[str, object], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总览"
    ws.append(["指标", "值"])
    ws.append(["总评测数", summary["total"]])
    ws.append(["通过数", summary["pass_count"]])
    ws.append(["问题数", summary["issue_count"]])
    ws.append(["综合准确率", summary["accuracy"]])
    ws.append(["综合问题率", summary["issue_rate"]])
    style_header(ws, 2)
    set_widths(ws, {"A": 22, "B": 18})

    ws = wb.create_sheet("维度问题")
    ws.append(["评测维度", "问题数", "问题率"])
    for dim, count in summary["dim_counts"].items():
        ws.append([dim, count, count / int(summary["total"])])
    style_header(ws, 3)
    set_widths(ws, {"A": 16, "B": 14, "C": 14})

    ws = wb.create_sheet("五级分类")
    ws.append(["五级分类", "总数", "通过数", "问题数", "准确率", "问题率", *DIMENSIONS])
    for l5, stat in sorted(summary["by_l5"].items(), key=lambda item: item[1]["accuracy"]):
        ws.append(
            [
                l5,
                stat["total"],
                stat["pass"],
                stat["issue"],
                stat["accuracy"],
                stat["issue_rate"],
                *[stat["dims"][dim] for dim in DIMENSIONS],
            ]
        )
    style_header(ws, 10)
    set_widths(ws, {"A": 16, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12})

    ws = wb.create_sheet("设备维度")
    ws.append(["query所属设备", "总数", "通过数", "问题数", "准确率", "问题率", *DIMENSIONS])
    for device, stat in summary["device_rank"]:
        ws.append(
            [
                device,
                stat["total"],
                stat["pass"],
                stat["issue"],
                stat["accuracy"],
                stat["issue_rate"],
                *[stat["dims"][dim] for dim in DIMENSIONS],
            ]
        )
    style_header(ws, 10)
    set_widths(ws, {"A": 18, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12})

    ws = wb.create_sheet("设备准确率分档")
    ws.append(["分档", "query所属设备", "总数", "问题数", "准确率", "问题率"])
    for band, items in summary["device_bands"].items():
        for device, stat in items:
            ws.append([band, device, stat["total"], stat["issue"], stat["accuracy"], stat["issue_rate"]])
    style_header(ws, 6)
    set_widths(ws, {"A": 14, "B": 34, "C": 12, "D": 12, "E": 12, "F": 12})

    ws = wb.create_sheet("五级维度矩阵")
    ws.append(["五级分类", *DIMENSIONS])
    for l5, stat in sorted(summary["by_l5"].items(), key=lambda item: item[1]["issue"], reverse=True):
        ws.append([l5, *[stat["dims"][dim] / stat["total"] for dim in DIMENSIONS]])
    style_header(ws, 5)
    set_widths(ws, {"A": 16, "B": 12, "C": 12, "D": 12, "E": 12})

    ws = wb.create_sheet("六级Top")
    ws.append(["五级分类", "六级分类", "总数", "问题数", "准确率", "主要问题维度"])
    for (l5, l6), stat in summary["top_l6"][:30]:
        top_dim = max(stat["dims"].items(), key=lambda item: item[1])[0]
        ws.append([l5, l6, stat["total"], stat["issue"], stat["accuracy"], top_dim])
    style_header(ws, 6)
    set_widths(ws, {"A": 16, "B": 18, "C": 12, "D": 12, "E": 12, "F": 16})

    ws = wb.create_sheet("典型案例")
    ws.append(["案例来源", "case_id", "query所属设备", "五级分类", "六级分类", "query", "DisplayText", "备注", *DIMENSIONS])
    case_groups = [
        ("最差设备", summary["bad_cases"]["devices"]),
        ("最差五级分类", summary["bad_cases"]["l5"]),
    ]
    for source, rows in case_groups:
        for row in rows:
            ws.append(
                [
                    source,
                    row["case_id"],
                    row["query所属设备"],
                    row["五级分类"],
                    row["六级分类"],
                    row["query"],
                    row["DisplayText"],
                    row["备注"],
                    *[row[dim] for dim in DIMENSIONS],
                ]
            )
    style_header(ws, 12)
    set_widths(ws, {"A": 14, "B": 18, "C": 18, "D": 14, "E": 16, "F": 44, "G": 56, "H": 34})

    ws = wb.create_sheet("历史典型案例")
    ws.append(["case_id", "query所属设备", "五级分类", "六级分类", "query", "DisplayText", "备注", *DIMENSIONS])
    for row in summary["examples"]:
        ws.append(
            [
                row["case_id"],
                row["query所属设备"],
                row["五级分类"],
                row["六级分类"],
                row["query"],
                row["DisplayText"],
                row["备注"],
                *[row[dim] for dim in DIMENSIONS],
            ]
        )
    style_header(ws, 11)
    set_widths(ws, {"A": 18, "B": 18, "C": 14, "D": 16, "E": 44, "F": 56, "G": 34})

    ws = wb.create_sheet("优先级")
    ws.append(["优先级", "参数", "重点设备", "问题数", "准确率", "主要问题", "建议动作"])
    for row in build_action_items(summary):
        ws.append(
            [
                row["priority"],
                row["param"],
                row["devices"],
                row["issue"],
                row["accuracy"],
                row["top_dim"],
                row["action"],
            ]
        )
    style_header(ws, 7)
    set_widths(ws, {"A": 12, "B": 22, "C": 24, "D": 12, "E": 12, "F": 14, "G": 54})

    ws = wb.create_sheet("低准确率设备失败参数")
    ws.append(["query所属设备", "总数", "问题数", "准确率", "主要问题维度", "失败参数"])
    for item in summary["low_accuracy_failures"]:
        params = "\n".join(f"{param} ({count})" for param, count in item["params"])
        ws.append([item["device"], item["total"], item["issue"], item["accuracy"], item["top_dim"], params])
    style_header(ws, 6)
    set_widths(ws, {"A": 34, "B": 12, "C": 12, "D": 12, "E": 14, "F": 56})

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.0%"
    wb.save(output_path)


def style_header(ws, cols: int) -> None:
    fill = PatternFill("solid", fgColor=COLORS["red"])
    font = Font(color=COLORS["white"], bold=True)
    for cell in ws[1][:cols]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def set_widths(ws, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_action_items(summary: dict[str, object]) -> list[dict[str, object]]:
    actions = []
    for (l5, l6), stat in summary["top_l6"][:7]:
        top_dim = max(stat["dims"].items(), key=lambda item: item[1])[0]
        device_hits = summary["by_l6_device"].get((l5, l6), [])
        devices = "\n".join(device for device, _ in device_hits) or "-"
        issue = int(stat["issue"])
        accuracy = float(stat["accuracy"])
        if issue >= 45 and (top_dim == "正确性" or accuracy < 0.80):
            priority = "P0"
        elif issue >= 35 or top_dim in {"正确性", "信息量"}:
            priority = "P1"
        else:
            priority = "P2"
        action = {
            "正确性": "补齐参数事实表，强化型号-参数精确匹配与单位校验",
            "信息量": "补充字段召回策略，约束答案必须给出完整规格值",
            "相关性": "优化规格查询意图识别，扩展参数别名和同义表达",
            "一致性": "增加答案前后一致性校验，统一支持/不支持表述",
        }[top_dim]
        actions.append(
            {
                "priority": priority,
                "param": f"{l5} / {l6}",
                "devices": devices,
                "issue": issue,
                "accuracy": accuracy,
                "top_dim": top_dim,
                "action": action,
            }
        )
    return actions


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int = 16,
    color: str = "dark",
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    auto_fit: bool = False,
) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    if auto_fit:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for idx, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(COLORS[color])
    return box


def add_rect(slide, x: float, y: float, w: float, h: float, fill: str, line: str | None = None) -> object:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(COLORS[fill])
    if line:
        shape.line.color.rgb = rgb(COLORS[line])
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def add_slide_header(slide, title: str, subtitle: str | None = None, label: str = "评测数据") -> None:
    add_text(slide, title, 0.55, 0.35, 8.2, 0.45, 21, "dark", True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 0.82, 8.7, 0.28, 9, "gray")
    add_rect(slide, 0.55, 1.12, 1.15, 0.035, "red")
    add_text(slide, label, 11.45, 0.42, 1.2, 0.28, 9, "red", True, PP_ALIGN.RIGHT)


def add_footer(slide, page: int) -> None:
    add_text(slide, f"{page:02d}", 12.55, 7.05, 0.35, 0.2, 8, "gray", False, PP_ALIGN.RIGHT)


def add_kpi_card(slide, x: float, y: float, w: float, h: float, label: str, value: str, sub: str, accent: str) -> None:
    add_rect(slide, x, y, w, h, "white", "light_gray")
    add_rect(slide, x, y, 0.06, h, accent)
    add_text(slide, label, x + 0.22, y + 0.18, w - 0.35, 0.25, 9, "gray", True)
    value_size = fitted_font_size(value, 25, fit_chars=12, min_size=7)
    add_text(slide, value, x + 0.22, y + 0.46, w - 0.35, 0.45, value_size, accent, True, auto_fit=True)
    add_text(slide, sub, x + 0.22, y + 1.03, w - 0.35, 0.25, 8, "gray")


def add_manual_bar_list(
    slide,
    items: list[tuple[str, float, str]],
    x: float,
    y: float,
    w: float,
    row_h: float,
    max_value: float = 1.0,
    value_fmt: str = "{:.1%}",
    label_w: float = 2.0,
    color: str = "red",
    limit: int | None = None,
    label_limit: int | None = 16,
    label_size: float = 8,
    label_fit_chars: int = 16,
) -> None:
    visible = items[:limit] if limit else items
    bar_x = x + label_w
    bar_w = w - label_w - 0.85
    for idx, (label, value, tag) in enumerate(visible):
        yy = y + idx * row_h
        display_label = str(label) if label_limit is None else truncate(label, label_limit)
        display_size = fitted_font_size(display_label, label_size, fit_chars=label_fit_chars, min_size=4.8)
        add_text(slide, display_label, x, yy + 0.02, label_w - 0.05, row_h - 0.02, display_size, "dark", auto_fit=True)
        add_rect(slide, bar_x, yy + 0.07, bar_w, row_h * 0.38, "light_gray")
        denominator = max_value or 1.0
        fill_w = max(0.02, bar_w * min(value / denominator, 1))
        add_rect(slide, bar_x, yy + 0.07, fill_w, row_h * 0.38, color)
        add_text(slide, value_fmt.format(value), bar_x + bar_w + 0.08, yy + 0.02, 0.75, row_h - 0.02, 8, "dark")
        if tag:
            add_text(slide, tag, bar_x + bar_w - 0.95, yy + 0.02, 0.9, row_h - 0.02, 7, "gray", False, PP_ALIGN.RIGHT)


def add_device_band_slide(
    prs: Presentation,
    blank,
    summary: dict[str, object],
    *,
    band: str,
    page: int,
    label: str,
) -> int:
    items = [
        (device, stat["accuracy"], f"n={stat['total']}")
        for device, stat in summary["device_bands"][band]
    ]
    chunks = [items[idx : idx + 32] for idx in range(0, len(items), 32)] or [[]]
    for chunk_idx, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(blank)
        add_rect(slide, 0, 0, 13.333, 7.5, "bg")
        suffix = f"（{chunk_idx + 1}/{len(chunks)}）" if len(chunks) > 1 else ""
        add_slide_header(slide, f"设备准确率分档：{band}{suffix}", "按答复准确率分档展示该区间内全部设备。", label)
        if not chunk:
            add_text(slide, f"当前没有准确率在 {band} 的设备。", 0.85, 2.35, 9.8, 0.5, 20, "gray", True)
            add_footer(slide, page)
            page += 1
            continue

        columns = split_evenly(chunk, 2)
        for col_idx, col_items in enumerate(columns):
            x = 0.65 + col_idx * 6.15
            add_manual_bar_list(
                slide,
                col_items,
                x,
                1.45,
                5.65,
                0.31,
                1.0,
                "{:.1%}",
                2.65,
                "red" if band == "0%-30%" else "orange" if band == "30%-60%" else "blue",
                label_limit=None,
                label_size=7.0,
                label_fit_chars=30,
            )
        add_footer(slide, page)
        page += 1
    return page


def split_evenly(items: list[tuple[str, float, str]], parts: int) -> list[list[tuple[str, float, str]]]:
    if not items:
        return [[]]
    size = (len(items) + parts - 1) // parts
    return [items[idx : idx + size] for idx in range(0, len(items), size)]


def add_l5_accuracy_slide(prs: Presentation, blank, summary: dict[str, object], *, page: int, label: str) -> int:
    items = [
        (l5, stat["accuracy"], f"n={stat['total']}")
        for l5, stat in sorted(summary["by_l5"].items(), key=lambda item: item[1]["accuracy"])
    ]
    chunks = [items[idx : idx + 28] for idx in range(0, len(items), 28)] or [[]]
    for chunk_idx, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(blank)
        add_rect(slide, 0, 0, 13.333, 7.5, "bg")
        suffix = f"（{chunk_idx + 1}/{len(chunks)}）" if len(chunks) > 1 else ""
        add_slide_header(slide, f"主要参数主体准确率{suffix}", "列举全部五级分类的答复准确率。", label)
        columns = split_evenly(chunk, 2)
        for col_idx, col_items in enumerate(columns):
            x = 0.8 + col_idx * 6.05
            add_manual_bar_list(
                slide,
                col_items,
                x,
                1.45,
                5.55,
                0.38,
                1.0,
                "{:.1%}",
                1.75,
                "blue",
                label_limit=None,
                label_size=8,
                label_fit_chars=20,
            )
        add_footer(slide, page)
        page += 1
    return page


def chunk_list(items: list[object], size: int) -> list[list[object]]:
    return [items[idx : idx + size] for idx in range(0, len(items), size)] or [[]]


def format_param_chunks(params: list[tuple[str, int]], *, max_items: int = 5, max_chars: int = 150) -> list[str]:
    if not params:
        return ["-"]
    chunks: list[str] = []
    current: list[str] = []
    current_chars = 0
    for param, count in params:
        line = f"{param} ({count})"
        if current and (len(current) >= max_items or current_chars + len(line) > max_chars):
            chunks.append("\n".join(current))
            current = []
            current_chars = 0
        current.append(line)
        current_chars += len(line)
    if current:
        chunks.append("\n".join(current))
    return chunks


def add_low_accuracy_failure_slides(prs: Presentation, blank, summary: dict[str, object], *, page: int, label: str) -> int:
    failures = summary["low_accuracy_failures"]
    expanded_rows: list[tuple[dict[str, object], str, int, int]] = []
    for item in failures:
        param_chunks = format_param_chunks(item["params"])
        for chunk_idx, params_text in enumerate(param_chunks):
            expanded_rows.append((item, params_text, chunk_idx, len(param_chunks)))

    row_groups = chunk_list(expanded_rows, 6)
    for group_idx, rows in enumerate(row_groups):
        slide = prs.slides.add_slide(blank)
        add_rect(slide, 0, 0, 13.333, 7.5, "bg")
        suffix = f"（{group_idx + 1}/{len(row_groups)}）" if len(row_groups) > 1 else ""
        add_slide_header(
            slide,
            f"0%-30% 低准确率机型失败参数{suffix}",
            "列举准确率 0%-30% 的全部机型；参数过多时同一机型会自动拆成多行。",
            label,
        )

        if not failures:
            add_text(slide, "当前没有准确率在 0%-30% 的机型。", 0.85, 2.45, 8.8, 0.45, 18, "green", True)
            add_footer(slide, page)
            page += 1
            continue

        x = 0.62
        y = 1.28
        widths = [2.45, 0.68, 0.72, 0.78, 0.95, 6.85]
        headers = ["机型", "总数", "问题数", "准确率", "主问题", "失败参数"]
        add_rect(slide, x, y, sum(widths), 0.38, "dark")
        xx = x
        for width, header in zip(widths, headers):
            add_text(slide, header, xx + 0.05, y + 0.08, width - 0.1, 0.18, 8, "white", True, PP_ALIGN.CENTER)
            xx += width

        row_h = 0.79
        for idx, (item, params_text, part_idx, part_count) in enumerate(rows):
            yy = y + 0.38 + idx * row_h
            add_rect(slide, x, yy, sum(widths), row_h, "white", "light_gray")
            xx = x
            device = str(item["device"])
            device_label = device if part_count == 1 else f"{device}\n参数 {part_idx + 1}/{part_count}"
            values = [
                device_label,
                str(item["total"]),
                str(item["issue"]),
                f"{item['accuracy']:.1%}",
                item["top_dim"],
                params_text,
            ]
            colors = ["dark", "dark", "dark", "red", "red", "dark"]
            for width, value, color, header in zip(widths, values, colors, headers):
                if header == "机型":
                    longest = max(str(value).splitlines() or [""], key=len)
                    font_size = fitted_font_size(longest, 7.2, fit_chars=26, min_size=5.0)
                    add_text(slide, str(value), xx + 0.05, yy + 0.05, width - 0.1, row_h - 0.1, font_size, color, auto_fit=True)
                elif header == "失败参数":
                    longest = max(str(value).splitlines() or [""], key=len)
                    font_size = fitted_font_size(longest, 6.5, fit_chars=48, min_size=4.8)
                    add_text(slide, str(value), xx + 0.05, yy + 0.05, width - 0.1, row_h - 0.1, font_size, color, auto_fit=True)
                else:
                    add_text(slide, str(value), xx + 0.06, yy + 0.22, width - 0.12, 0.22, 8, color, header in {"准确率", "主问题"}, PP_ALIGN.CENTER)
                xx += width

        add_text(
            slide,
            "说明：失败参数为该机型综合评测结果=1 的「五级分类 / 六级分类」，括号内为失败 case 数。",
            0.75,
            6.55,
            10.6,
            0.25,
            9,
            "gray",
        )
        add_footer(slide, page)
        page += 1
    return page


def heat_color(rate: float, max_rate: float) -> str:
    if rate == 0:
        return "DCFCE7"
    if max_rate <= 0:
        return "DCFCE7"
    ratio = min(rate / max_rate, 1.0)
    if ratio < 0.25:
        return "FEE2E2"
    if ratio < 0.5:
        return "FCA5A5"
    if ratio < 0.75:
        return "F87171"
    return "C7000B"


def add_heatmap(slide, summary: dict[str, object], x: float, y: float, w: float, h: float) -> None:
    by_l5 = summary["by_l5"]
    l5_rows = sorted(by_l5.items(), key=lambda item: item[1]["issue_rate"], reverse=True)
    if not l5_rows:
        add_text(slide, "暂无五级分类数据。", x, y + 1.8, w, 0.4, 18, "gray", True, PP_ALIGN.CENTER)
        return
    label_w = 2.1
    cell_w = (w - label_w) / len(DIMENSIONS)
    row_h = h / (len(l5_rows) + 1)
    add_rect(slide, x, y, w, row_h, "dark")
    add_text(slide, "五级分类", x + 0.06, y + 0.05, label_w - 0.12, row_h - 0.05, 8, "white", True)
    for col, dim in enumerate(DIMENSIONS):
        add_text(slide, dim, x + label_w + col * cell_w, y + 0.05, cell_w, row_h - 0.05, 8, "white", True, PP_ALIGN.CENTER)

    max_rate = max(stat["dims"][dim] / stat["total"] for _, stat in l5_rows for dim in DIMENSIONS)
    for row_idx, (l5, stat) in enumerate(l5_rows, start=1):
        yy = y + row_idx * row_h
        add_rect(slide, x, yy, label_w, row_h, "white", "light_gray")
        add_text(slide, truncate(l5, 12), x + 0.06, yy + 0.05, label_w - 0.12, row_h - 0.04, 8, "dark")
        for col, dim in enumerate(DIMENSIONS):
            rate = stat["dims"][dim] / stat["total"]
            color_hex = heat_color(rate, max_rate)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x + label_w + col * cell_w),
                Inches(yy),
                Inches(cell_w),
                Inches(row_h),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(color_hex)
            shape.line.color.rgb = rgb(COLORS["white"])
            add_text(
                slide,
                f"{rate:.1%}",
                x + label_w + col * cell_w,
                yy + 0.05,
                cell_w,
                row_h - 0.04,
                8,
                "white" if max_rate > 0 and rate / max_rate >= 0.75 else "dark",
                False,
                PP_ALIGN.CENTER,
            )


def create_ppt(
    summary: dict[str, object],
    *,
    output_path: Path,
    product_name: str = "手机",
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    label = "评测数据"

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, "bg")
    add_rect(slide, 0, 0, 13.333, 0.18, "red")
    add_text(slide, f"小艺{product_name}全量参数\n知识覆盖评测报告", 0.85, 1.45, 7.2, 1.5, 34, "dark", True)
    cover_subtitle = f"华为商城在售{product_name}全系列规格查询"
    add_text(slide, cover_subtitle, 0.9, 3.25, 7.2, 0.35, 14, "gray")
    add_kpi_card(slide, 8.8, 1.55, 3.25, 1.35, "样本规模", f"{summary['total']:,}", f"覆盖 {len(summary['by_device'])} 个设备", "red")
    add_kpi_card(slide, 8.8, 3.15, 3.25, 1.35, "综合准确率", f"{summary['accuracy']:.1%}", "人工综合评测结果=0", "green")
    add_kpi_card(slide, 8.8, 4.75, 3.25, 1.35, "主要问题", str(summary["top_dim"]), "四维度中命中最多", "orange")
    add_text(slide, f"生成日期：{date.today().isoformat()}", 0.9, 6.7, 3.0, 0.25, 9, "gray")
    add_footer(slide, 1)

    # 2. Summary dashboard
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, "bg")
    add_slide_header(slide, "综合看板", "合并展示核心结论、整体准确率和四维度问题分布。", label)
    add_kpi_card(slide, 0.6, 1.28, 2.45, 1.08, "总评测数", f"{summary['total']:,}", f"设备 {len(summary['by_device'])} 个", "red")
    add_kpi_card(slide, 3.22, 1.28, 2.45, 1.08, "通过数", f"{summary['pass_count']:,}", "无问题 case", "green")
    add_kpi_card(slide, 5.84, 1.28, 2.45, 1.08, "问题数", f"{summary['issue_count']:,}", "至少一个维度命中", "orange")
    add_kpi_card(slide, 8.46, 1.28, 2.45, 1.08, "综合准确率", f"{summary['accuracy']:.1%}", "通过数 / 总评测数", "blue")
    worst_dim_device = summary["worst_dim_device"]
    add_kpi_card(slide, 11.08, 1.28, 1.6, 1.08, "最差维度设备", str(worst_dim_device["device"]), f"{worst_dim_device['dimension']}维度", "red")

    chart_data = ChartData()
    chart_data.categories = ["通过", "问题"]
    chart_data.add_series("综合结果", (summary["pass_count"], summary["issue_count"]))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.82), Inches(2.65), Inches(3.2), Inches(3.2), chart_data).chart
    chart.has_legend = False
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = "0%"
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.BEST_FIT
    chart.series[0].points[0].format.fill.solid()
    chart.series[0].points[0].format.fill.fore_color.rgb = rgb(COLORS["green"])
    chart.series[0].points[1].format.fill.solid()
    chart.series[0].points[1].format.fill.fore_color.rgb = rgb(COLORS["red"])
    add_text(slide, f"{summary['accuracy']:.1%}", 1.48, 3.74, 1.85, 0.4, 23, "green", True, PP_ALIGN.CENTER)
    add_text(slide, "综合准确率", 1.58, 4.2, 1.65, 0.22, 8, "gray", False, PP_ALIGN.CENTER)

    dim_items = [
        (dim, summary["dim_counts"][dim] / summary["total"], f"{summary['dim_counts'][dim]:,}条")
        for dim in DIMENSIONS
    ]
    add_text(slide, "四维度问题率", 4.45, 2.72, 2.2, 0.25, 12, "dark", True)
    add_manual_bar_list(slide, dim_items, 4.45, 3.1, 4.1, 0.48, max(v for _, v, _ in dim_items), "{:.1%}", 1.05, "red")

    insights = [
        f"整体准确率 {summary['accuracy']:.1%}，问题 case {summary['issue_count']:,} 条。",
        f"主要短板集中在「{summary['top_dim']}」，应优先处理该维度的高频参数。",
        f"最差维度为「{summary['top_dim']}」，该维度最差设备为「{worst_dim_device['device']}」（问题率 {worst_dim_device['dimension_rate']:.1%}）。",
    ]
    for idx, text in enumerate(insights):
        y = 3.0 + idx * 0.72
        add_rect(slide, 9.15, y, 3.35, 0.5, "white", "light_gray")
        add_rect(slide, 9.15, y, 0.055, 0.5, "red" if idx == 0 else "orange")
        add_text(slide, text, 9.35, y + 0.12, 2.85, 0.22, 9, "dark", idx == 0)
    add_text(slide, "判断方式：综合结果=1 表示四个维度中至少一个存在问题；四维度命中率可重叠，不做互斥占比。", 0.88, 6.38, 11.6, 0.25, 8, "gray")
    add_footer(slide, 2)

    page = 3
    if summary["device_bands"].get("90%-100%"):
        page = add_device_band_slide(prs, blank, summary, band="90%-100%", page=page, label=label)
    page = add_device_band_slide(prs, blank, summary, band="60%-90%", page=page, label=label)
    page = add_device_band_slide(prs, blank, summary, band="30%-60%", page=page, label=label)
    page = add_device_band_slide(prs, blank, summary, band="0%-30%", page=page, label=label)
    page = add_l5_accuracy_slide(prs, blank, summary, page=page, label=label)

    # Heatmap
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, "bg")
    add_slide_header(slide, "五级分类 × 问题维度矩阵", "颜色越深，表示该分类在该评测维度上的问题率越高。", label)
    add_heatmap(slide, summary, 0.8, 1.35, 11.65, 5.4)
    add_footer(slide, page)
    page += 1

    # Top L6
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, "bg")
    add_slide_header(slide, "六级分类问题 Top 20", "按具体参数的问题 case 数排序，右侧灰字为累计覆盖的问题占比。", label)
    top20 = summary["top_l6"][:20]
    total_issue = summary["issue_count"]
    cumulative = 0
    items = []
    max_issue = max((stat["issue"] for _, stat in top20), default=0)
    for (_l5, l6), stat in top20:
        cumulative += stat["issue"]
        label = f"{l6}"
        tag = f"{cumulative / total_issue:.0%}" if total_issue else "0%"
        items.append((label, stat["issue"], tag))
    if items and max_issue:
        add_manual_bar_list(slide, items, 0.75, 1.28, 8.25, 0.255, max_issue, "{:.0f}", 1.55, "red")
    else:
        add_text(slide, "当前没有问题 case，暂无六级分类 Top 问题。", 0.9, 2.3, 7.2, 0.4, 18, "green", True)
    add_text(slide, "读法：红色条形长度代表失败 case 数，条末数字是该参数失败 case 数；右侧灰字是从第1名累加到当前行覆盖全部问题 case 的比例。", 0.82, 6.55, 8.4, 0.35, 9, "gray")
    add_rect(slide, 9.45, 1.35, 3.05, 4.9, "white", "light_gray")
    top_coverage = sum(stat["issue"] for _, stat in top20) / total_issue if total_issue else 0
    add_text(slide, "Top 20 覆盖", 9.75, 1.75, 1.4, 0.3, 10, "gray", True)
    add_text(slide, f"{top_coverage:.1%}", 9.75, 2.12, 1.7, 0.45, 28, "red", True)
    add_text(slide, "的问题 case", 11.25, 2.28, 0.9, 0.25, 10, "gray")
    top_param = top20[0] if top20 else None
    add_text(slide, "最高频参数", 9.75, 3.05, 1.4, 0.3, 10, "gray", True)
    if top_param and top_param[1]["issue"]:
        add_text(slide, f"{top_param[0][0]} / {top_param[0][1]}", 9.75, 3.42, 2.35, 0.45, 15, "dark", True)
        add_text(slide, f"问题数 {top_param[1]['issue']}，准确率 {top_param[1]['accuracy']:.1%}", 9.75, 4.02, 2.25, 0.25, 9, "gray")
        top_devices = "、".join(device for device, _ in summary["by_l6_device"].get(top_param[0], []))
        add_text(slide, f"高发设备：{top_devices or '-'}", 9.75, 4.38, 2.25, 0.25, 9, "gray")
    else:
        add_text(slide, "暂无", 9.75, 3.42, 2.35, 0.45, 15, "green", True)
    add_text(slide, "用途：优先修复 Top 参数，通常能最快降低整体问题数。", 9.75, 5.15, 2.25, 0.55, 11, "dark")
    add_footer(slide, page)
    page += 1

    # Examples
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, "bg")
    add_slide_header(slide, "典型错误案例", "2个最差设备 badcase + 2个最差五级分类 badcase。", label)
    positions = [(0.75, 1.35), (6.8, 1.35), (0.75, 4.05), (6.8, 4.05)]
    example_rows = summary["bad_cases"]["devices"] + summary["bad_cases"]["l5"]
    example_titles = ["最差设备", "最差设备", "最差五级分类", "最差五级分类"]
    for idx, (row, (x, y), case_title) in enumerate(zip(example_rows, positions, example_titles), start=1):
        add_rect(slide, x, y, 5.65, 2.25, "white", "light_gray")
        dims = " / ".join([dim for dim in DIMENSIONS if row[dim]])
        add_text(slide, f"{case_title} {idx}｜{row['query所属设备']}｜{row['五级分类']} / {row['六级分类']}", x + 0.2, y + 0.14, 4.15, 0.25, 9, "red", True)
        add_text(slide, dims, x + 4.1, y + 0.18, 1.25, 0.25, 9, "orange", True, PP_ALIGN.RIGHT)
        add_text(slide, f"Q：{row['query']}", x + 0.2, y + 0.46, 5.15, 0.45, fitted_font_size(row["query"], 8.0, fit_chars=58, min_size=5.6), "dark", True, auto_fit=True)
        add_text(slide, f"A：{row['DisplayText']}", x + 0.2, y + 0.97, 5.15, 0.52, fitted_font_size(row["DisplayText"], 7.5, fit_chars=70, min_size=5.3), "gray", auto_fit=True)
        add_text(slide, f"备注：{row['备注']}", x + 0.2, y + 1.6, 5.15, 0.45, fitted_font_size(row["备注"], 8.0, fit_chars=58, min_size=5.6), "dark", auto_fit=True)
    add_footer(slide, page)
    page += 1

    page = add_low_accuracy_failure_slides(prs, blank, summary, page=page, label=label)

    prs.save(output_path)


def main() -> None:
    product_name = str(CONFIG["PRODUCT_NAME"])
    rows = load_eval_rows(Path(CONFIG["INPUT_XLSX"]))
    summary = compute_summary(rows)
    pptx_path = Path(CONFIG["OUTPUT_PPTX"])
    summary_path = pptx_path.with_name(f"{pptx_path.stem}_统计.xlsx")
    write_summary_workbook(summary, summary_path)
    print(f"summary={summary_path}")
    create_ppt(
        summary,
        output_path=pptx_path,
        product_name=product_name,
    )
    print(f"pptx={pptx_path}")
    print(f"total={summary['total']} accuracy={summary['accuracy']:.2%} issue={summary['issue_count']}")


if __name__ == "__main__":
    main()
