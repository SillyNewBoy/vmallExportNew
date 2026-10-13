"""Export category JSON payloads to Excel workbooks and CSV files."""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

FLAT_COLUMNS = [
    "category_key",
    "category_name",
    "series_name",
    "product_name",
    "product_name_short",
    "group_name",
    "attr_name",
    "attr_value",
    "attr_remark",
    "product_url",
    "specs_url",
]


def flatten_payload(payload: dict[str, Any]) -> list[dict[str, Any]]:
    """Flatten a category JSON tree into one row per specification item."""
    rows: list[dict[str, Any]] = []
    for series in (payload.get("series") or {}).values():
        for product in (series.get("products") or {}).values():
            base = {
                "category_key": payload.get("category_key", ""),
                "category_name": payload.get("category_name", ""),
                "series_name": series.get("series_name", ""),
                "product_name": product.get("product_name", ""),
                "product_name_short": product.get("product_name_short", ""),
                "product_url": product.get("product_url", ""),
                "specs_url": product.get("specs_url", ""),
            }
            specs = product.get("specifications") or []
            if not specs:
                rows.append(
                    {
                        **base,
                        "group_name": "",
                        "attr_name": "",
                        "attr_value": "",
                        "attr_remark": "",
                    }
                )
                continue
            for spec in specs:
                rows.append(
                    {
                        **base,
                        "group_name": spec.get("group_name", ""),
                        "attr_name": spec.get("attr_name", ""),
                        "attr_value": spec.get("attr_value", ""),
                        "attr_remark": spec.get("attr_remark", ""),
                    }
                )
    return rows


def write_csv(payload: dict[str, Any], output_dir: str | Path) -> Path:
    """Write the flattened payload to ``{category_key}.csv``."""
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    file_path = output_path / f"{payload.get('category_key') or 'category'}.csv"
    rows = flatten_payload(payload)
    with file_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=FLAT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote CSV export: %s (%s rows)", file_path, len(rows))
    return file_path


def write_excel(payload: dict[str, Any], output_dir: str | Path) -> Path:
    """Write the payload to ``{category_key}.xlsx`` with one sheet per series.

    The first sheet ``全部参数`` always contains every specification row.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    file_path = output_path / f"{payload.get('category_key') or 'category'}.xlsx"

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="C7000B")

    def fill_sheet(sheet, rows: list[dict[str, Any]]) -> None:
        sheet.append(FLAT_COLUMNS)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            sheet.append([row.get(column, "") for column in FLAT_COLUMNS])
        sheet.freeze_panes = "A2"
        for idx, column in enumerate(FLAT_COLUMNS, start=1):
            width = max((len(str(row.get(column, ""))) for row in rows[:200]), default=8)
            sheet.column_dimensions[get_column_letter(idx)].width = min(
                max(width + 2, len(column) + 2), 60
            )

    all_rows = flatten_payload(payload)
    fill_sheet(workbook.active, all_rows)
    workbook.active.title = "全部参数"

    for series in (payload.get("series") or {}).values():
        series_rows = flatten_payload({"series": {series.get("series_key"): series}})
        sheet_name = _safe_sheet_name(str(series.get("series_name") or "未分类"))
        if sheet_name in workbook.sheetnames:
            suffix = 2
            while f"{sheet_name}_{suffix}" in workbook.sheetnames:
                suffix += 1
            sheet_name = f"{sheet_name}_{suffix}"
        fill_sheet(workbook.create_sheet(sheet_name), series_rows)

    workbook.save(file_path)
    logger.info("Wrote Excel export: %s (%s rows)", file_path, len(all_rows))
    return file_path


def _safe_sheet_name(name: str) -> str:
    for char in '[]:*?/\\':
        name = name.replace(char, " ")
    return name.strip()[:31] or "sheet"
